# export_to_docx.py

import docx
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from converts.ast_nodes import *
from exports.settings import ExportSettings
from exports.text_utils import remove_emojis
from exports.docx_spaces import add_run_preserved, preserve_run_spaces


class ExportToDocx:

    def __init__(self, settings: Optional[ExportSettings] = None):
        self.settings = settings or ExportSettings()
        self.doc = docx.Document()
        self.list_stack = []  # Stack for tracking nested lists

    def export(self, ast: Document) -> Document:
        for node in ast.children:
            self._export_node(node)

        return self.doc

    def _export_node(self, node: ASTNode):
        if isinstance(node, Heading):
            self._export_heading(node)
        elif isinstance(node, Paragraph):
            self._export_paragraph(node)
        elif isinstance(node, CodeBlock):
            self._export_code_block(node)
        elif isinstance(node, List):
            self._export_list(node)
        elif isinstance(node, BlockQuote):
            self._export_blockquote(node)
        elif isinstance(node, HorizontalRule):
            if not self.settings.strip_horizontal_rules:
                self._export_horizontal_rule()
        elif isinstance(node, Table):
            self._export_table(node)

    def _export_heading(self, node: Heading):
        heading = self.doc.add_heading('', level=node.level)
        self._fill_paragraph_with_inline(heading, node.content)

    def _export_paragraph(self, node: Paragraph):
        paragraph = self.doc.add_paragraph()
        self._fill_paragraph_with_inline(paragraph, node.content)

    def _export_code_block(self, node: CodeBlock):
        lines = node.code.split('\n')
        if not lines:
            lines = ['']

        for line in lines:
            paragraph = self.doc.add_paragraph()
            text = self._clean_text(line)
            run = add_run_preserved(paragraph, text)

            if self.settings.plain_code:
                self._apply_run_style(run)
            else:
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(100, 100, 100)
                self._set_paragraph_background(paragraph, 'E8E8E8')
                paragraph.paragraph_format.left_indent = Inches(0.3)
                paragraph.paragraph_format.right_indent = Inches(0.3)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)

    def _export_list(self, node: List):
        for item in node.items:
            if node.ordered:
                paragraph = self.doc.add_paragraph(style='List Number')
            else:
                paragraph = self.doc.add_paragraph(style='List Bullet')

            self._fill_paragraph_with_inline(paragraph, item.content)

    def _export_blockquote(self, node: BlockQuote):
        if self.settings.strip_blockquotes:
            for child in node.content:
                self._export_node(child)
            return

        for child in node.content:
            self._export_node(child)
            # Accepting quote style for the last paragraph
            if self.doc.paragraphs:
                last_paragraph = self.doc.paragraphs[-1]
                last_paragraph.paragraph_format.left_indent = Inches(0.5)
                last_paragraph.paragraph_format.right_indent = Inches(0.5)

                # Adding left border
                self._set_paragraph_border(last_paragraph, left_color='4472C4', left_width=24)

                # Gray background
                self._set_paragraph_background(last_paragraph, 'F2F2F2')

    def _export_horizontal_rule(self):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)

        # Adding a horizontal line through the paragraph's border
        pPr = paragraph._element.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '24')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _export_table(self, node: Table):
        if self.settings.table_mode == 'word_native':
            self._export_table_word_native(node)
        else:
            self._export_table_styled(node)

    def _export_table_styled(self, node: Table):
        num_cols = len(node.header)
        num_rows = len(node.rows) + 1
        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.style = 'Light Grid Accent 1'

        header_cells = table.rows[0].cells
        for i, header_text in enumerate(node.header):
            cell = header_cells[i]
            cell.text = ''
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(self._clean_text(header_text))
            preserve_run_spaces(run)
            if not self.settings.strip_bold:
                run.font.bold = True
            self._apply_run_style(run, in_table=True)
            self._set_paragraph_background(paragraph, 'D3D3D3')

        for row_idx, row_data in enumerate(node.rows, start=1):
            row_cells = table.rows[row_idx].cells
            for col_idx, cell_content in enumerate(row_data):
                # Checking borders
                if col_idx >= num_cols:
                    break
                cell = row_cells[col_idx]
                cell.text = ''
                paragraph = cell.paragraphs[0]
                self._fill_paragraph_with_inline(paragraph, cell_content, in_table=True)

    def _export_table_word_native(self, node: Table):
        num_cols = max(len(node.header), 1)
        num_rows = len(node.rows) + 1

        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        self._set_table_borders(table)

        for i, header_text in enumerate(node.header):
            cell = table.rows[0].cells[i]
            self._set_cell_plain_text(cell, header_text, header=True)

        for row_idx, row_data in enumerate(node.rows, start=1):
            for col_idx, cell_content in enumerate(row_data):
                if col_idx >= num_cols:
                    break
                cell = table.rows[row_idx].cells[col_idx]
                text = self._clean_text(self._extract_text(cell_content))
                self._set_cell_plain_text(cell, text, header=False)

    def _set_cell_plain_text(self, cell, text: str, header: bool = False):
        cell.text = ''
        paragraph = cell.paragraphs[0]
        run = add_run_preserved(paragraph, text)
        if header and not self.settings.strip_bold:
            run.font.bold = True
        self._apply_run_style(run, in_table=True)

    def _set_table_borders(self, table):
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        if tbl_pr is None:
            tbl_pr = OxmlElement('w:tblPr')
            tbl.insert(0, tbl_pr)

        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement(f'w:{edge}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), '000000')
            borders.append(border)
        tbl_pr.append(borders)

    def _fill_paragraph_with_inline(
        self,
        paragraph,
        inline_nodes: List[ASTNode],
        in_table: bool = False,
    ):
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)

        self._add_inline_nodes(paragraph, inline_nodes, in_table=in_table)

    def _clean_text(self, text: str) -> str:
        if self.settings.strip_emojis:
            return remove_emojis(text)
        return text

    def _apply_run_style(self, run, in_table: bool = False):
        run.font.color.rgb = self.settings.table_rgb if in_table else self.settings.rgb
        font_name = self.settings.font_name
        if font_name:
            run.font.name = font_name

    def _add_inline_nodes(
        self,
        paragraph,
        inline_nodes: List[ASTNode],
        bold: bool = False,
        italic: bool = False,
        in_table: bool = False,
    ):
        effective_bold = bold and not self.settings.strip_bold
        effective_italic = italic and not self.settings.strip_italic

        for node in inline_nodes:
            if isinstance(node, Text):
                run = add_run_preserved(paragraph, self._clean_text(node.value))
                run.bold = effective_bold
                run.italic = effective_italic
                self._apply_run_style(run, in_table=in_table)

            elif isinstance(node, Bold):
                if self.settings.strip_bold:
                    self._add_inline_nodes(
                        paragraph, node.content, bold=bold, italic=italic, in_table=in_table,
                    )
                else:
                    self._add_inline_nodes(
                        paragraph, node.content, bold=True, italic=italic, in_table=in_table,
                    )

            elif isinstance(node, Italic):
                if self.settings.strip_italic:
                    self._add_inline_nodes(
                        paragraph, node.content, bold=bold, italic=italic, in_table=in_table,
                    )
                else:
                    self._add_inline_nodes(
                        paragraph, node.content, bold=bold, italic=True, in_table=in_table,
                    )

            elif isinstance(node, Code):
                if self.settings.strip_backticks:
                    self._add_inline_nodes(
                        paragraph,
                        [Text(value=self._clean_text(node.value))],
                        bold=bold, italic=italic, in_table=in_table,
                    )
                else:
                    run = add_run_preserved(paragraph, self._clean_text(node.value))  # ← Добавлен _clean_text
                    run.font.name = 'Courier New'
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(200, 0, 0)

            elif isinstance(node, Link):
                self._add_hyperlink(paragraph, node.url, node.text)

            elif isinstance(node, Image):
                try:
                    paragraph.add_run().add_picture(node.url, width=Inches(4))
                except Exception:
                    run = add_run_preserved(paragraph, f"[Image: {node.alt}]")
                    if not self.settings.strip_italic:
                        run.italic = True
                    self._apply_run_style(run, in_table=in_table)

    def _add_hyperlink(self, paragraph, url: str, text_nodes: List[ASTNode]):
        try:
            part = paragraph.part
            rel_id = part.relate_to(
                url,
                'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                is_external=True
            )

            r = OxmlElement('w:r')
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), rel_id)

            run_element = OxmlElement('w:r')
            text_content = self._clean_text(self._extract_text(text_nodes))
            t = OxmlElement('w:t')
            t.set(qn('xml:space'), 'preserve')
            t.text = text_content
            run_element.append(t)
            hyperlink.append(run_element)
            paragraph._element.append(hyperlink)
        except Exception:
            # If creating a link fails, add it as text
            text_content = self._clean_text(self._extract_text(text_nodes))
            run = add_run_preserved(paragraph, f"{text_content} ({url})")
            self._apply_run_style(run)

    def _extract_text(self, nodes: List[ASTNode]) -> str:
        text = ''
        for node in nodes:
            if isinstance(node, Text):
                text += node.value
            elif isinstance(node, Bold) or isinstance(node, Italic) or isinstance(node, Code):
                text += self._extract_text(
                    node.content if hasattr(node, 'content') else [Text(value=node.value)])
        return text

    def _set_paragraph_background(self, paragraph, color: str):
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), color)
        paragraph._element.get_or_add_pPr().append(shading_elm)

    def _set_paragraph_border(self, paragraph, left_color: str = None, left_width: int = 12):
        pPr = paragraph._element.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')

        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), str(left_width))
        left.set(qn('w:space'), '0')
        left.set(qn('w:color'), left_color or '000000')

        pBdr.append(left)
        pPr.append(pBdr)

    def save(self, filename: str):
        self.doc.save(filename)
